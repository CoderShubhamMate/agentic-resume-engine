#!/usr/bin/env python3
"""
jobmatch.py - Resume vs Job Description matcher (Jobscan-style) -> PDF report

Usage:
    python jobmatch.py resume.pdf jd.txt -o report.pdf [--title "Senior Backend Engineer"]
    (resume: .pdf/.docx/.txt   |   jd: file path OR raw text)

Install:  pip install pdfplumber python-docx reportlab
"""
import argparse, datetime, math, re
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

# ============================================================
# 1. TEXT EXTRACTION
# ============================================================
def read_text(src):
    p = Path(src)
    if not p.exists():                       # treat as raw pasted text
        return str(src)
    ext = p.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(p) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    if ext == ".docx":
        import docx
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    return p.read_text(errors="ignore")

# ============================================================
# 2. KNOWLEDGE BASE  (extend these; in production use ESCO / O*NET / a DB)
# ============================================================
SKILLS = {  # canonical: [aliases]
    "python": [], "java": [], "javascript": ["js", "ecmascript"], "typescript": ["ts"],
    "c++": ["cpp"], "c#": ["csharp"], "golang": [], "rust": [], "php": [], "ruby": [],
    "sql": [], "nosql": [], "postgresql": ["postgres"], "mysql": [], "mongodb": ["mongo"],
    "redis": [], "elasticsearch": ["elastic search"], "react": ["reactjs", "react.js"],
    "angular": [], "vue": ["vuejs", "vue.js"], "node.js": ["nodejs"], "django": [],
    "flask": [], "fastapi": [], "spring boot": ["springboot", "spring"], "html": [], "css": [],
    "aws": ["amazon web services"], "azure": [], "gcp": ["google cloud"], "docker": [],
    "kubernetes": ["k8s"], "terraform": [], "ansible": [], "jenkins": [],
    "ci/cd": ["cicd", "continuous integration", "continuous delivery"], "git": [],
    "linux": [], "rest api": ["restful", "rest apis"], "graphql": [], "grpc": [],
    "microservices": ["microservice"], "system design": [], "machine learning": ["ml"],
    "deep learning": [], "nlp": ["natural language processing"], "llm": ["llms", "large language models"],
    "pytorch": [], "tensorflow": [], "scikit-learn": ["sklearn"], "pandas": [], "numpy": [],
    "spark": ["pyspark"], "kafka": [], "airflow": [], "etl": [], "data analysis": ["data analytics"],
    "data modeling": ["data modelling"], "tableau": [], "power bi": ["powerbi"], "excel": [],
    "agile": ["scrum"], "jira": [], "unit testing": ["unit tests", "pytest", "junit"],
    "seo": [], "google analytics": [], "salesforce": [], "figma": [], "product management": [],
}
SOFT = {
    "communication": ["communicate"], "leadership": ["led", "lead", "leading"],
    "collaboration": ["collaborate", "collaborative"], "problem solving": ["problem-solving"],
    "teamwork": ["team player"], "stakeholder management": ["stakeholders"],
    "mentoring": ["mentor", "mentored"], "ownership": [], "adaptability": ["adaptable"],
    "critical thinking": [], "time management": [], "decision making": ["decision-making"],
}
STOP = set("""a an the and or of to in for on with as at by from is are be been being this that these
those it its you your we our they their will can may must should would could have has had do does did not
no nor but if then than so such into about over under per via etc including include includes within across
using use used work working role team ability able strong excellent good great plus years year experience
experienced looking join company candidate candidates responsibilities requirements qualifications preferred
required skills skill knowledge understanding job position opportunity benefits salary apply who what when
where how all any more most other some also well new""".split())

# ============================================================
# 3. NLP HELPERS
# ============================================================
def stem(w):
    for suf in ("ing", "ed", "es", "s"):
        if w.endswith(suf) and len(w) - len(suf) >= 4:
            return w[: -len(suf)]
    return w

STOP_STEMS = STOP | {stem(w) for w in STOP}
tokens = lambda t: re.findall(r"[a-z][a-z0-9+#]{1,}", t.lower())

def term_pattern(names):
    alts = "|".join(re.escape(n) for n in sorted(names, key=len, reverse=True))
    return re.compile(rf"(?<![\w+#.])(?:{alts})(?![\w+#])", re.I)

def jd_weighted_lines(jd):
    """Tag each JD line with importance: Requirements=1.5, Preferred=0.6, other=1.0."""
    out, cur = [], 1.0
    for line in jd.splitlines():
        s = line.strip().lower()
        heading = s.endswith(":") or line.strip().isupper() or s.startswith("#")
        if heading and len(s) < 70:
            if re.search(r"requirement|qualification|must|required|minimum|you.ll need|what we.re looking", s):
                cur = 1.5
            elif re.search(r"preferred|nice to have|bonus|good to have|a plus", s):
                cur = 0.6
            else:
                cur = 1.0
        out.append((line, cur))
    return out

# ============================================================
# 4. KEYWORD EXTRACTION + MATCHING
# ============================================================
def analyse_dictionary(jd_lines, resume, vocab):
    rows = []
    for canon, aliases in vocab.items():
        pat = term_pattern([canon] + aliases)
        freq, secw = 0, 0.0
        for line, w in jd_lines:
            n = len(pat.findall(line))
            if n:
                freq += n
                secw = max(secw, w)
        if not freq:
            continue
        rc = len(pat.findall(resume))
        rows.append(dict(name=canon, jd_count=freq, resume_count=rc, found=rc > 0,
                         required=secw >= 1.5, nice=secw < 1.0, weight=secw * (1 + 0.5 * min(freq - 1, 3))))
    return sorted(rows, key=lambda d: (-d["weight"], d["name"]))

def extract_generic_keywords(jd, resume, skip_words, top=15):
    raw = tokens(jd)
    st = [stem(t) for t in raw]
    surface = {}
    for r, s in zip(raw, st):
        surface.setdefault(s, Counter())[r] += 1
    ok = lambda s: s not in STOP_STEMS and len(s) >= 3 and s not in skip_words
    uni, bi = Counter(), Counter()
    for i, s in enumerate(st):
        if not ok(s):
            continue
        uni[s] += 1
        if i + 1 < len(st) and ok(st[i + 1]):
            bi[(s, st[i + 1])] += 1
    cands = [(k, c * 2) for k, c in bi.items() if c >= 2]
    used = {w for k, _ in cands for w in k}
    cands += [((k,), c) for k, c in uni.items() if c >= 2 and k not in used]
    cands.sort(key=lambda x: -x[1])
    rtext = " ".join(stem(t) for t in tokens(resume))
    rset = set(rtext.split())
    out = []
    for k, score in cands[:top]:
        name = " ".join(surface[w].most_common(1)[0][0] for w in k)
        if len(k) == 1:
            credit = 1.0 if k[0] in rset else 0.0
        elif re.search(r"\b" + r"\s".join(map(re.escape, k)) + r"\b", rtext):
            credit = 1.0
        else:
            credit = 0.5 if all(w in rset for w in k) else 0.0
        out.append(dict(name=name, weight=score, credit=credit, found=credit > 0))
    return out

# ============================================================
# 5. OTHER SIGNALS
# ============================================================
def title_score(title, resume):
    tt = [stem(t) for t in tokens(title) if stem(t) not in STOP_STEMS]
    if not tt:
        return None
    rtext = " ".join(stem(t) for t in tokens(resume))
    if " ".join(tt) in rtext:
        return 1.0
    return 0.8 * sum(t in rtext.split() for t in tt) / len(tt)

def resume_years(text):
    MONTH_MAP = {
        "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
        "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6,
        "jul": 7, "july": 7, "aug": 8, "august": 8, "sep": 9, "september": 9, "sept": 9,
        "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12
    }
    today = datetime.date.today()
    now_m = today.year * 12 + today.month

    pat = (
        r"\b(?:([A-Za-z]{3,9})\.?\s+)?((?:19|20)\d{2})\s*(?:-|\u2013|\u2014|to)\s*"
        r"(?:(?:([A-Za-z]{3,9})\.?\s+)?((?:19|20)\d{2})|(present|current|now))\b"
    )

    spans = []
    for m in re.finditer(pat, text, re.I):
        m1_str, y1_str, m2_str, y2_str, pres = m.groups()
        y1 = int(y1_str)
        m1 = MONTH_MAP.get(m1_str.lower() if m1_str else "", 1)
        start_idx = y1 * 12 + (m1 - 1)

        if pres:
            end_idx = now_m
        else:
            y2 = int(y2_str)
            m2 = MONTH_MAP.get(m2_str.lower() if m2_str else "", 12)
            end_idx = y2 * 12 + m2

        if start_idx <= end_idx:
            spans.append((start_idx, end_idx))

    if not spans:
        return 0.0

    spans.sort()
    total_months = 0
    end_bound = None
    for a, b in spans:
        if end_bound is None or a > end_bound:
            total_months += (b - a)
            end_bound = b
        elif b > end_bound:
            total_months += (b - end_bound)
            end_bound = b

    return round(total_months / 12.0, 1)

def required_years(jd):
    m = re.search(r"(\d{1,2})\+?\s*(?:-\s*\d{1,2}\s*)?(?:years|yrs)", jd, re.I)
    return int(m.group(1)) if m else None

def cosine(a, b):
    ca = Counter(stem(t) for t in tokens(a) if stem(t) not in STOP_STEMS)
    cb = Counter(stem(t) for t in tokens(b) if stem(t) not in STOP_STEMS)
    dot = sum(ca[k] * cb[k] for k in ca)
    na, nb = math.sqrt(sum(v * v for v in ca.values())), math.sqrt(sum(v * v for v in cb.values()))
    return dot / (na * nb) if na and nb else 0.0

def ats_checks(resume):
    words = len(resume.split())
    has = lambda p: bool(re.search(p, resume, re.I))
    return [
        ("Email address present", has(r"[\w.+-]+@[\w-]+\.[\w.]+")),
        ("Phone number present", has(r"\+?\d[\d\s().-]{8,}\d")),
        ("Experience section found", has(r"\b(experience|employment|work history)\b")),
        ("Education section found", has(r"\b(education|university|college|b\.?tech|b\.?sc|m\.?sc|degree)\b")),
        ("Skills section found", has(r"\bskills\b")),
        ("Length between 300 and 1200 words (%d now)" % words, 300 <= words <= 1200),
        ("3+ quantified achievements (%, $, numbers)",
         len(re.findall(r"\d+(?:\.\d+)?\s*%|\$\s?\d|\b\d+(?:\.\d+)?\s?(?:x|k|m|million|billion|users|"
                        r"customers|clients|ms|hrs|hours)\b", resume, re.I)) >= 3),
    ]

# ============================================================
# 6. SCORING  (weights sum to 100; missing components are re-normalised)
# ============================================================
WEIGHTS = {"Hard skills": 40, "Job keywords": 20, "Job title": 10, "Text similarity": 10,
           "Experience years": 8, "Soft skills": 5, "ATS format": 7}

def analyse(resume, jd, title=None):
    lines = jd_weighted_lines(jd)
    title = title or next((l.strip() for l in jd.splitlines() if l.strip()), "")
    hard = analyse_dictionary(lines, resume, SKILLS)
    soft = analyse_dictionary(lines, resume, SOFT)
    skip = {stem(w) for r in hard + soft for w in tokens(r["name"])}
    kws = extract_generic_keywords(jd, resume, skip)
    checks = ats_checks(resume)
    ry, need = resume_years(resume), required_years(jd)

    wsum = lambda rows, f: sum(r["weight"] * f(r) for r in rows) / sum(r["weight"] for r in rows) if rows else None
    comp = {
        "Hard skills": wsum(hard, lambda r: r["found"]),
        "Job keywords": wsum(kws, lambda r: r["credit"]),
        "Job title": title_score(title, resume),
        "Text similarity": min(1.0, cosine(resume, jd) / 0.45),
        "Experience years": None if need is None else min(1.0, ry / need) if need else 1.0,
        "Soft skills": (sum(r["found"] for r in soft) / len(soft)) if soft else None,
        "ATS format": sum(ok for _, ok in checks) / len(checks),
    }
    live = {k: v for k, v in comp.items() if v is not None}
    overall = 100 * sum(WEIGHTS[k] * v for k, v in live.items()) / sum(WEIGHTS[k] for k in live)
    return dict(title=title, overall=round(overall), comp=comp, hard=hard, soft=soft, kws=kws,
                checks=checks, resume_years=ry, need_years=need)

def recommendations(r):
    recs = []
    miss = [h for h in r["hard"] if not h["found"]]
    req = [h["name"] for h in miss if h["required"]]
    opt = [h["name"] for h in miss if not h["required"]]
    if req:
        recs.append("Add these REQUIRED skills if you genuinely have them (Skills section AND inside an "
                    "experience bullet showing how you used them): " + ", ".join(req) + ".")
    if opt:
        recs.append("Nice-to-have skills missing: " + ", ".join(opt[:8]) + ". Add only those that are true for you.")
    kmiss = [k["name"] for k in r["kws"] if not k["found"]]
    if kmiss:
        recs.append("Mirror the job's wording where accurate. Missing terms: " + ", ".join(kmiss[:8]) + ".")
    ts = r["comp"]["Job title"]
    if ts is not None and ts < 1:
        recs.append("Put the target job title ('%s') in your headline/summary so ATS title matching succeeds." % r["title"])
    if r["need_years"] and r["resume_years"] < r["need_years"]:
        recs.append("JD asks for %d+ years; your dated roles total about %d. State total experience clearly in "
                    "the summary and make sure every role has a start-end date." % (r["need_years"], r["resume_years"]))
    sm = [s["name"] for s in r["soft"] if not s["found"]]
    if sm:
        recs.append("Show soft skills through results, not adjectives (JD stresses: " + ", ".join(sm) + ").")
    for label, ok in r["checks"]:
        if not ok:
            recs.append("Format/ATS fix: " + label + ".")
    if not recs:
        recs.append("Strong match. Fine-tune by moving the most relevant skills to the top third of page 1.")
    return recs

# ============================================================
# 7. PDF REPORT
# ============================================================
def build_pdf(r, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Drawing, Rect
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    GREEN, AMBER, RED = colors.HexColor("#1b873f"), colors.HexColor("#c77700"), colors.HexColor("#c62828")
    NAVY, LIGHT = colors.HexColor("#1f2d4d"), colors.HexColor("#f2f4f8")
    col = lambda s: GREEN if s >= .8 else AMBER if s >= .6 else RED
    ss = getSampleStyleSheet()
    H = ParagraphStyle("H", parent=ss["Heading2"], textColor=NAVY, spaceBefore=12, spaceAfter=6)
    N = ParagraphStyle("N", parent=ss["Normal"], fontSize=9.5, leading=13)

    def bar(s):
        d = Drawing(110, 9)
        d.add(Rect(0, 0, 100, 8, fillColor=colors.HexColor("#dde1e8"), strokeColor=None))
        d.add(Rect(0, 0, 100 * s, 8, fillColor=col(s), strokeColor=None))
        return d

    def table(data, widths, status_col=None):
        t = Table(data, colWidths=widths, repeatRows=1)
        st = [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#c9ced8")),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        if status_col is not None:
            for i, row in enumerate(data[1:], 1):
                ok = row[status_col].startswith("Found") or row[status_col] == "OK" or row[status_col].startswith("Partial")
                bad = row[status_col] in ("MISSING", "FIX")
                st += [("TEXTCOLOR", (status_col, i), (status_col, i), RED if bad else GREEN if ok else AMBER),
                       ("FONTNAME", (status_col, i), (status_col, i), "Helvetica-Bold")]
        t.setStyle(TableStyle(st))
        return t

    ov = r["overall"]
    verdict = "Strong match" if ov >= 80 else "Moderate match - improve before applying" if ov >= 60 else "Low match - needs significant tailoring"
    story = [Paragraph("Resume - Job Match Report", ss["Title"]),
             Paragraph("Target role: <b>%s</b> &nbsp;|&nbsp; Generated: %s" % (escape(r["title"]), datetime.date.today()), N),
             Spacer(1, 10)]
    BIG = ParagraphStyle("BIG", parent=N, fontSize=34, leading=40, alignment=1)
    big = Table([[Paragraph('<font color="#%s"><b>%d%%</b></font>' % (col(ov / 100).hexval()[2:], ov), BIG),
                  Paragraph("<b>%s</b><br/>%d of %d hard skills from the JD found in your resume."
                            % (verdict, sum(h["found"] for h in r["hard"]), len(r["hard"])), N)]],
                colWidths=[40 * mm, 130 * mm])
    big.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .8, NAVY), ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 10),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    story += [big, Paragraph("Score breakdown", H)]

    rows = [["Component", "Weight", "Score", ""]]
    for k, v in r["comp"].items():
        rows.append([k, "%d" % WEIGHTS[k], "n/a (not in JD)" if v is None else "%d%%" % round(v * 100), "" if v is None else bar(v)])
    story.append(table(rows, [55 * mm, 20 * mm, 35 * mm, 60 * mm]))

    story.append(Paragraph("Hard skills (from job description)", H))
    rows = [["Skill", "Priority", "In JD", "In resume", "Status"]]
    for h in sorted(r["hard"], key=lambda x: x["found"]):
        rows.append([h["name"], "Required" if h["required"] else "Nice-to-have" if h["nice"] else "Normal", h["jd_count"], h["resume_count"],
                     "Found" if h["found"] else "MISSING"])
    story.append(table(rows, [55 * mm, 30 * mm, 20 * mm, 25 * mm, 40 * mm], status_col=4) if r["hard"]
                 else Paragraph("No known hard skills detected in the JD.", N))

    if r.get("verified_omitted"):
        story.append(Paragraph("In your Master_Data but missing from this resume (safe to add)", H))
        rows = [["Skill", "Priority", "Evidence"]]
        for item in r["verified_omitted"]:
            ev = Paragraph(escape(str(item.get("evidence", ""))), N)
            rows.append([item.get("name", ""), item.get("priority", "Normal"), ev])
        story.append(table(rows, [45 * mm, 30 * mm, 95 * mm]))

    if r.get("true_gaps"):
        story.append(Paragraph("Not supported by your Master_Data (do NOT add)", H))
        rows = [["Skill", "Priority"]]
        for item in r["true_gaps"]:
            rows.append([item.get("name", ""), item.get("priority", "Normal")])
        story.append(table(rows, [100 * mm, 70 * mm]))

    story.append(Paragraph("Other important job keywords", H))
    rows = [["Keyword", "Status"]] + [[k["name"], "Found" if k["credit"] == 1 else "Partial" if k["credit"] else "MISSING"] for k in r["kws"]]
    story.append(table(rows, [100 * mm, 70 * mm], status_col=1) if r["kws"] else Paragraph("None extracted.", N))

    if r["soft"]:
        story.append(Paragraph("Soft skills", H))
        rows = [["Skill", "Status"]] + [[s["name"], "Found" if s["found"] else "MISSING"] for s in r["soft"]]
        story.append(table(rows, [100 * mm, 70 * mm], status_col=1))

    story.append(Paragraph("ATS and format checks", H))
    rows = [["Check", "Status"]] + [[l, "OK" if ok else "FIX"] for l, ok in r["checks"]]
    story.append(table(rows, [130 * mm, 40 * mm], status_col=1))

    if r.get("unverified_in_resume"):
        WARN = colors.HexColor("#c62828")
        story.append(Paragraph("Skills in resume NOT supported by Master_Data — REMOVE", H))
        rows = [["Skill"]]
        for u in r["unverified_in_resume"]:
            rows.append([Paragraph(f"<font color='#c62828'><b>{escape(u)}</b></font>", N)])
        story.append(table(rows, [170 * mm]))

    if r.get("reasons"):
        story.append(Paragraph("Reasons this score is not higher", H))
        for reason in r["reasons"]:
            story.append(Paragraph("• " + escape(reason), N))
            story.append(Spacer(1, 3))

    story.append(Paragraph("What to improve (prioritised)", H))
    for i, rec in enumerate(recommendations(r), 1):
        story.append(Paragraph("<b>%d.</b> %s" % (i, escape(rec)), N))
        story.append(Spacer(1, 4))
    story += [Spacer(1, 10), Paragraph("<i>Only add keywords that reflect real experience. "
              "Keyword stuffing is flagged by recruiters and hurts you in interviews.</i>", N)]
    SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm, title="Resume Match Report").build(story)


# ============================================================
# 8. CLI
# ============================================================
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("resume"); ap.add_argument("jd")
    ap.add_argument("-o", "--out", default="match_report.pdf")
    ap.add_argument("--title", help="Job title (default: first line of the JD)")
    a = ap.parse_args()
    result = analyse(read_text(a.resume), read_text(a.jd), a.title)
    build_pdf(result, a.out)
    print("Match score: %d%%  ->  %s" % (result["overall"], a.out))
