import openpyxl, os, sys, datetime, re, argparse
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PATH_PORTFOLIO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'Application List.xlsx'))

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

ROW_FONT = Font(name="Calibri", size=10, color="111827")
ALT_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

STATUS_STYLES = {
    "Applied": (PatternFill(start_color="E8F1F5", end_color="E8F1F5", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="1B4F72")),
    "Rejected": (PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="78281F")),
    "Scam": (PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="7E5109")),
    "Selected / Offer": (PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="196F3D")),
    "Accepted / Offer": (PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="196F3D")),
    "Interview Scheduled": (PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="512E5F")),
    "In Discussion": (PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid"), Font(name="Calibri", size=10, bold=True, color="512E5F")),
}

THIN_BORDER = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

headers = ["Company", "Role", "Date", "Status", "Platform / Source", "Email Thread / Link"]

def load_existing(file_path):
    if not os.path.exists(file_path):
        return []
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=False):
        if r[0].value is not None:
            c = r[0].value
            role = r[1].value
            d = r[2].value
            s = r[3].value
            src = r[4].value
            e_title = r[5].value
            e_url = r[5].hyperlink.target if r[5].hyperlink else None
            
            if isinstance(d, datetime.datetime):
                d = d.date()
                
            rows.append({
                "company": c,
                "role": role,
                "date": d,
                "status": s,
                "source": src,
                "email_title": e_title,
                "email_url": e_url
            })
    return rows

def write_excel(records, file_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications Master"
    ws.views.sheetView[0].showGridLines = True

    # Header
    ws.row_dimensions[1].height = 26
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style='medium', color='0F2A4A'))

    for row_idx, item in enumerate(records, 2):
        ws.row_dimensions[row_idx].height = 20
        row_fill = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
        
        c_cell = ws.cell(row=row_idx, column=1, value=item["company"])
        c_cell.font = Font(name="Calibri", size=10, bold=True, color="111827")
        c_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        r_cell = ws.cell(row=row_idx, column=2, value=item["role"])
        r_cell.font = ROW_FONT
        r_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        d_val = item["date"]
        d_cell = ws.cell(row=row_idx, column=3, value=d_val)
        d_cell.font = ROW_FONT
        d_cell.number_format = 'yyyy-mm-dd'
        d_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        stat_val = item["status"]
        s_cell = ws.cell(row=row_idx, column=4, value=stat_val)
        s_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        s_fill, s_font = STATUS_STYLES.get(stat_val, (row_fill, ROW_FONT))
        s_cell.fill = s_fill
        s_cell.font = s_font
        
        p_cell = ws.cell(row=row_idx, column=5, value=item.get("source", "Portfolio Resume"))
        p_cell.font = ROW_FONT
        p_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        e_title = item.get("email_title")
        e_url = item.get("email_url")
        e_cell = ws.cell(row=row_idx, column=6)
        if e_title:
            e_cell.value = e_title
            if e_url:
                e_cell.hyperlink = e_url
                e_cell.font = Font(name="Calibri", size=10, color="0551C4", underline="single")
            else:
                e_cell.font = ROW_FONT
        else:
            e_cell.font = ROW_FONT
        e_cell.alignment = Alignment(horizontal="left", vertical="center")

        for col_num in range(1, 7):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = THIN_BORDER
            if col_num != 4:
                cell.fill = row_fill

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format == 'yyyy-mm-dd' and isinstance(cell.value, (datetime.date, datetime.datetime)):
                val_str = cell.value.strftime('%Y-%m-%d')
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ensure parent dir exists
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    wb.save(file_path)

def add_application(company, role, date_str=None, status="Applied", source="Portfolio Resume", link_title=None, link_url=None):
    if not date_str:
        dt = datetime.date.today()
    else:
        dt = datetime.datetime.strptime(date_str, "%Y-%m-%d").date()

    existing = load_existing(PATH_PORTFOLIO)
    
    # check if key already exists
    key = (company.strip().lower(), role.strip().lower())
    found = False
    for item in existing:
        ik = (item["company"].strip().lower(), item["role"].strip().lower())
        if ik == key or (len(company) > 3 and company.strip().lower() in ik[0]):
            item["date"] = dt
            item["status"] = status
            item["source"] = source
            if link_title: item["email_title"] = link_title
            if link_url: item["email_url"] = link_url
            found = True
            break

    if not found:
        existing.append({
            "company": company.strip(),
            "role": role.strip(),
            "date": dt,
            "status": status,
            "source": source,
            "email_title": link_title,
            "email_url": link_url
        })

    # Sort newest first
    existing.sort(key=lambda x: (x["date"] if isinstance(x["date"], datetime.date) else datetime.date(2020,1,1)), reverse=True)

    write_excel(existing, PATH_PORTFOLIO)
    print(f"Successfully logged application: '{company}' - '{role}' on {dt.strftime('%Y-%m-%d')} as '{status}'")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Log job application to Excel file")
    parser.add_argument("--company", required=True, help="Company Name")
    parser.add_argument("--role", required=True, help="Job Role")
    parser.add_argument("--date", help="Application date (YYYY-MM-DD)")
    parser.add_argument("--status", default="Applied", help="Status (Applied, Rejected, Scam, etc.)")
    parser.add_argument("--source", default="Portfolio Resume", help="Source platform")
    parser.add_argument("--link-title", help="Email thread link title")
    parser.add_argument("--link-url", help="Email thread link URL")
    
    args = parser.parse_args()
    add_application(args.company, args.role, args.date, args.status, args.source, args.link_title, args.link_url)
